# -*- coding: utf-8 -*-
"""读取 G系列MODBUS协议IO点表 xlsx，导出为 UTF-8 文本"""
import openpyxl

import os
path = r"D:\work\副本G系列MODBUS协议IO点表20250322.xlsx"
out = os.path.join(os.environ.get("TEMP", r"C:\Windows\Temp"), "protocol_xlsx_dump.txt")

wb = openpyxl.load_workbook(path, data_only=True)
lines = []
lines.append("Sheets: %s" % wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    lines.append("")
    lines.append("### Sheet: %s  max_row=%d max_col=%d" % (name, ws.max_row, ws.max_column))
    lines.append("Merged: %s" % [str(r) for r in ws.merged_cells.ranges])
    lines.append("")
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True), start=1):
        vals = []
        for v in row:
            if v is None:
                vals.append("")
            else:
                vals.append(str(v).replace("\n", "\\n"))
        lines.append("R%d: %s" % (r_idx, " | ".join(vals)))

with open(out, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("written:", out, "lines:", len(lines))
