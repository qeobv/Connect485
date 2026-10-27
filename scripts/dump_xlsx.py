# -*- coding: utf-8 -*-
"""通用 xlsx 结构/内容 dump 工具：输出到 TEMP 下 UTF-8 文本"""
import os
import sys
import openpyxl

path = sys.argv[1] if len(sys.argv) > 1 else r"D:\work\多回转S1-Z系列 MODBUS协议IO点表20250104(1) - 副本.xlsx"
out = os.path.join(os.environ.get("TEMP", r"C:\Windows\Temp"), "s1z_dump.txt")

wb = openpyxl.load_workbook(path, data_only=True)
lines = []
lines.append("Sheets: %s" % wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    lines.append("")
    lines.append("### Sheet: %s  max_row=%d max_col=%d" % (name, ws.max_row, ws.max_column))
    lines.append("Merged: %s" % [str(r) for r in ws.merged_cells.ranges])
    lines.append("")
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True), start=1):
        vals = []
        for v in row:
            if v is None:
                vals.append("")
            else:
                vals.append(str(v).replace("\n", "\\n"))
        lines.append("R%d: %s" % (r_idx, " | ".join(vals)))

with open(out, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("written:", out, "lines:", len(lines))
